#!/usr/bin/env python3
"""Generate and validate the multisubject catalog from the general workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

from catalog_core import validate_catalog


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "data" / "catalogo-actividades.xlsx"
DEFAULT_CATALOG = ROOT / "data" / "generated" / "catalog.json"
DEFAULT_REPORT = ROOT / "data" / "generated" / "generation-report.json"
DELIMITER = "|"


def clean(value):
    return value.strip() if isinstance(value, str) else value


def rows_as_dicts(sheet):
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [clean(value) for value in values[0]]
    rows = []
    for row_number, values_row in enumerate(values[1:], start=2):
        if not any(value is not None for value in values_row):
            continue
        record = {headers[index]: clean(value) for index, value in enumerate(values_row)}
        record["__row__"] = row_number
        rows.append(record)
    return rows


def split_values(value) -> list[str]:
    if value in (None, ""):
        return []
    result = []
    seen = set()
    for item in str(value).split(DELIMITER):
        normalized = item.strip()
        key = normalized.casefold()
        if normalized and key not in seen:
            result.append(normalized)
            seen.add(key)
    return result


def iso_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def optional(value):
    return None if value in (None, "") else value


def without_empty(mapping: dict) -> dict:
    return {key: value for key, value in mapping.items() if value not in (None, "", [], {})}


def load_general_workbook(source: Path) -> tuple[dict, dict[str, dict]]:
    workbook = openpyxl.load_workbook(source, data_only=False)
    required = {
        "ASIGNATURAS",
        "TEMAS",
        "MIGRACIÓN_TEMAS",
        "TIPOS_ACTIVIDAD",
        "PLATAFORMAS",
        "ACTIVIDADES",
        "ACTIVIDAD_TEMA",
        "CRÉDITOS",
        "CONTROL_DE_CALIDAD",
        "INFORMACIÓN",
    }
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise SystemExit(f"Faltan hojas obligatorias: {', '.join(sorted(missing))}")

    raw = {name: rows_as_dicts(workbook[name]) for name in required}
    locations: dict[str, dict] = {}

    subjects = []
    for row in raw["ASIGNATURAS"]:
        subject_id = row["ID"]
        locations[subject_id] = {"sheet": "ASIGNATURAS", "row": row["__row__"]}
        visual = without_empty({
            "accent": row["Color principal"],
            "accentSoft": row["Color suave"],
            "icon": row["Icono"],
            "image": row["Imagen"],
            "imageAlt": row["Texto alternativo"],
            "imageObjectPosition": row["Posición de imagen"],
            "creditId": row["ID de crédito"],
        })
        subjects.append(without_empty({
            "id": subject_id,
            "name": row["Nombre"],
            "shortName": row["Nombre corto"],
            "slug": row["Slug"],
            "legacySlugs": split_values(row["Slugs históricos"]),
            "order": int(row["Orden"]),
            "description": row["Descripción"],
            "visual": visual,
            "status": row["Estado"],
        }))

    topics = []
    for row in raw["TEMAS"]:
        topic_id = row["ID"]
        locations[topic_id] = {"sheet": "TEMAS", "row": row["__row__"]}
        visual = without_empty({
            "image": row["Imagen"],
            "imageAlt": row["Texto alternativo"],
            "imageObjectPosition": row["Posición de imagen"],
            "creditId": row["ID de crédito"],
        })
        legacy = without_empty({
            "sourceId": row["ID legado"],
            "originalName": row["Nombre original"],
            "historicalCourse": row["Curso histórico"],
            "declaredCount": int(row["Conteo declarado legado"]) if row["Conteo declarado legado"] not in (None, "") else None,
        })
        topics.append(without_empty({
            "id": topic_id,
            "subjectId": row["ID de asignatura"],
            "name": row["Nombre"],
            "slug": row["Slug"],
            "legacySlugs": split_values(row["Slugs históricos"]),
            "order": int(row["Orden"]),
            "description": row["Descripción"],
            "visual": visual,
            "status": row["Estado"],
            "legacy": legacy,
        }))

    activity_types = []
    for row in raw["TIPOS_ACTIVIDAD"]:
        type_id = row["ID"]
        locations[type_id] = {"sheet": "TIPOS_ACTIVIDAD", "row": row["__row__"]}
        activity_types.append(without_empty({
            "id": type_id,
            "name": row["Nombre"],
            "icon": row["Icono"],
            "order": int(row["Orden"]),
        }))

    platforms = []
    for row in raw["PLATAFORMAS"]:
        platform_id = row["ID"]
        locations[platform_id] = {"sheet": "PLATAFORMAS", "row": row["__row__"]}
        platforms.append(without_empty({
            "id": platform_id,
            "name": row["Nombre"],
            "hostnames": split_values(row["Hosts"]),
            "color": row["Color"],
            "order": int(row["Orden"]),
        }))

    credits = []
    for row in raw["CRÉDITOS"]:
        credit_id = row["ID"]
        locations[credit_id] = {"sheet": "CRÉDITOS", "row": row["__row__"]}
        credits.append(without_empty({
            "id": credit_id,
            "title": row["Título"],
            "author": row["Autoría"],
            "sourceName": row["Fuente"],
            "sourceUrl": row["URL de fuente"],
            "license": row["Licencia"],
            "licenseUrl": row["URL de licencia"],
            "note": row["Nota"],
        }))

    topic_ids_by_activity: dict[str, list[str]] = defaultdict(list)
    for row in raw["ACTIVIDAD_TEMA"]:
        topic_ids_by_activity[row["ID de actividad"]].append(row["ID de tema"])

    activities = []
    for row in raw["ACTIVIDADES"]:
        activity_id = row["ID"]
        locations[activity_id] = {"sheet": "ACTIVIDADES", "row": row["__row__"]}
        source_kind = row["Clase de origen"]
        if source_kind == "external":
            source_data = without_empty({
                "kind": "external",
                "platformId": row["ID de plataforma"],
                "resourceId": str(row["ID de recurso"]) if row["ID de recurso"] not in (None, "") else None,
                "url": row["URL pública"],
                "canonicalUrl": row["URL canónica"],
                "originalUrl": row["URL original"],
                "linkStatus": row["Estado del enlace"],
                "lastVerifiedAt": iso_date(row["Última verificación"]),
                "titleVerified": bool(row["Título verificado"]),
                "titleSource": row["Fuente del título"],
                "descriptionSource": row["Fuente de la descripción"],
                "verificationNote": row["Notas de verificación"],
            })
        elif source_kind == "native":
            source_data = {
                "kind": "native",
                "native": {
                    "engineVersion": row["Versión de motor nativo"],
                    "activityType": row["Tipo de actividad nativa"],
                    "contentPath": row["Ruta de contenido nativo"],
                },
            }
        else:
            source_data = {"kind": source_kind}
        topic_ids = topic_ids_by_activity.get(activity_id, [])
        activities.append(without_empty({
            "schemaVersion": 1,
            "id": activity_id,
            "slug": row["Slug"],
            "legacySlugs": split_values(row["Slugs históricos"]),
            "title": row["Título"],
            "sourceTitle": row["Título de origen"],
            "description": row["Descripción"],
            "primarySubjectId": row["ID de asignatura principal"],
            "subjectIds": split_values(row["IDs de asignatura"]),
            "primaryTopicId": optional(row["ID de tema principal"]),
            "topicIds": topic_ids,
            "language": row["Idioma"],
            "typeId": row["ID de tipo"],
            "originCourseLabel": row["Curso histórico"],
            "tags": split_values(row["Etiquetas"]),
            "keywords": split_values(row["Palabras clave"]),
            "dates": {
                "createdAt": iso_date(row["Fecha de creación"]),
                "publishedAt": iso_date(row["Fecha de publicación"]),
                "updatedAt": iso_date(row["Fecha de actualización"]),
            },
            "publicationStatus": row["Estado editorial"],
            "source": source_data,
            "legacy": without_empty({
                "subjectLabel": row["Asignatura original"],
                "linkStatusLabel": row["Estado de enlace original"],
            }),
        }))

    quality_issues = []
    for row in raw["CONTROL_DE_CALIDAD"]:
        quality_issues.append({
            "activityId": row["ID de actividad"],
            "url": row["URL"],
            "issueType": row["Tipo de incidencia"],
            "explanation": row["Explicación"],
            "recommendedAction": row["Acción recomendada"],
            "status": row["Estado"],
        })

    information = [{"field": row["Campo"], "explanation": row["Explicación"]} for row in raw["INFORMACIÓN"]]
    information_by_field = {item["field"]: item["explanation"] for item in information}
    topic_id_map = {row["ID legado de sección"]: row["Nuevo ID de tema"] for row in raw["MIGRACIÓN_TEMAS"]}
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    catalog = {
        "schemaVersion": 1,
        "source": {"file": "data/catalogo-actividades.xlsx", "sha256": source_hash},
        "subjects": sorted(subjects, key=lambda item: (item["order"], item["id"])),
        "topics": sorted(topics, key=lambda item: (item["subjectId"], item["order"], item["id"])),
        "activityTypes": sorted(activity_types, key=lambda item: (item["order"], item["id"])),
        "platforms": sorted(platforms, key=lambda item: (item["order"], item["id"])),
        "activities": activities,
        "credits": credits,
        "qualityIssues": quality_issues,
        "information": information,
        "migration": {
            "legacySource": information_by_field.get("Fuente histórica"),
            "legacySourceSha256": information_by_field.get("SHA-256 de la fuente histórica"),
            "topicIdMap": topic_id_map,
        },
    }
    return catalog, locations


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def print_problems(label: str, items: list[dict]) -> None:
    if not items:
        return
    print(f"{label} ({len(items)}):")
    for item in items:
        location = f"{item.get('sheet')} fila {item.get('row') or '?'}"
        if item.get("column"):
            location += f", columna {item['column']}"
        print(f"- [{item.get('code')}] {location}; ID {item.get('id')}: {item.get('problem')}. Esperado: {item.get('expected')}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    catalog, locations = load_general_workbook(args.input)
    errors, warnings = validate_catalog(catalog, locations)
    if errors:
        print_problems("Errores bloqueantes", errors)
        raise SystemExit(1)

    activities = catalog["activities"]
    platform_counts = Counter(activity["source"].get("platformId") for activity in activities if activity["source"]["kind"] == "external")
    language_counts = Counter(activity["language"] for activity in activities)
    source_kind_counts = Counter(activity["source"]["kind"] for activity in activities)
    multi_topic_ids = [activity["id"] for activity in activities if len(activity["topicIds"]) > 1]
    missing_created_ids = [activity["id"] for activity in activities if activity["dates"]["createdAt"] is None]
    if missing_created_ids:
        warnings.insert(0, {
            "code": "missing-created-at",
            "sheet": "ACTIVIDADES",
            "row": None,
            "column": "Fecha de creación",
            "id": None,
            "problem": f"{len(missing_created_ids)} actividades no tienen fecha original de creación",
            "expected": "Mantener null hasta disponer de una fuente fiable",
            "activityIds": missing_created_ids,
        })

    report = {
        "schemaVersion": 1,
        "source": catalog["source"],
        "counts": {
            "subjects": len(catalog["subjects"]),
            "topics": len(catalog["topics"]),
            "activities": len(activities),
            "relations": sum(len(activity["topicIds"]) for activity in activities),
            "multiTopicActivities": len(multi_topic_ids),
        },
        "byPlatform": dict(sorted(platform_counts.items())),
        "byLanguage": dict(sorted(language_counts.items())),
        "bySourceKind": dict(sorted(source_kind_counts.items())),
        "multiTopicActivityIds": multi_topic_ids,
        "warnings": warnings,
        "errors": [],
        "migration": catalog["migration"],
    }
    write_json(args.catalog, catalog)
    write_json(args.report, report)
    print(json.dumps({**report["counts"], "byPlatform": report["byPlatform"], "byLanguage": report["byLanguage"]}, ensure_ascii=False, indent=2))
    print_problems("Avisos", warnings)


if __name__ == "__main__":
    main()
