#!/usr/bin/env python3
"""Append one reviewed external activity to the source workbook, atomically."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import unicodedata
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from copy import copy

from catalog_core import BCP47, validate_catalog
from importlib.machinery import SourceFileLoader


ROOT = Path(__file__).resolve().parents[1]
generator = SourceFileLoader("catalog_generator", str(ROOT / "scripts/generate-catalog.py")).load_module()
DEFAULT_BOOK = ROOT / "data/catalogo-actividades.xlsx"


class ImportErrorMessage(ValueError):
    pass


def required_string(payload, name, limit):
    value = payload.get(name)
    if not isinstance(value, str) or not value.strip() or len(value) > limit or any(ord(ch) < 32 for ch in value):
        raise ImportErrorMessage(f"{name}: introduce un texto válido de hasta {limit} caracteres")
    return value.strip()


def slugify(title, resource_id):
    normalized = unicodedata.normalize("NFKD", title)
    ascii_title = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    base = re.sub(r"[^a-z0-9]+", "-", ascii_title.lower()).strip("-")[:75].rstrip("-")
    return f"{base or 'actividad'}-{resource_id}"


def validate_input(payload, catalog):
    if not isinstance(payload, dict) or set(payload) - {"url", "platformId", "resourceId", "title", "description", "subjectId", "topicIds", "typeId", "language", "verified", "sourceTitle", "manual"}:
        raise ImportErrorMessage("La solicitud contiene campos inesperados")
    url = required_string(payload, "url", 2048)
    platform = required_string(payload, "platformId", 40)
    resource = required_string(payload, "resourceId", 24)
    title = required_string(payload, "title", 180)
    description = required_string(payload, "description", 1200)
    subject = required_string(payload, "subjectId", 80)
    kind = required_string(payload, "typeId", 80)
    language = required_string(payload, "language", 35)
    topics = payload.get("topicIds")
    if not isinstance(topics, list) or not 1 <= len(topics) <= 20 or not all(isinstance(t, str) for t in topics) or len(set(topics)) != len(topics):
        raise ImportErrorMessage("Selecciona uno o más temas sin repetirlos")
    if payload.get("verified") not in (True, False):
        raise ImportErrorMessage("verified debe ser booleano")
    source_title = payload.get("sourceTitle")
    if source_title is not None and (not isinstance(source_title, str) or len(source_title) > 180):
        raise ImportErrorMessage("sourceTitle inválido")
    if not re.fullmatch(r"[0-9]{1,20}", resource):
        raise ImportErrorMessage("No se pudo obtener un ID externo válido")
    if not BCP47.fullmatch(language):
        raise ImportErrorMessage("Idioma BCP 47 inválido")
    if subject not in {s["id"] for s in catalog["subjects"] if s["status"] == "active"}:
        raise ImportErrorMessage("La asignatura seleccionada ya no está activa")
    topic_map = {t["id"]: t for t in catalog["topics"]}
    if any(t not in topic_map or topic_map[t]["subjectId"] != subject or topic_map[t]["status"] != "active" for t in topics):
        raise ImportErrorMessage("Algún tema ya no existe o no pertenece a la asignatura")
    if kind not in {t["id"] for t in catalog["activityTypes"]}:
        raise ImportErrorMessage("El tipo de actividad no existe")
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if parsed.scheme != "https" or parsed.username or parsed.password or parsed.port or parsed.fragment or parsed.query:
        raise ImportErrorMessage("La URL pública debe ser HTTPS y no contener credenciales, query ni fragmento")
    if platform == "wordwall" and host == "wordwall.net":
        match = re.fullmatch(r"/[a-z]{2}/resource/([0-9]+)(?:/[^?#]*)?", parsed.path)
        prefix = "WW"
    elif platform == "educaplay" and host in {"es.educaplay.com", "www.educaplay.com"}:
        match = re.fullmatch(r"/(?:recursos-educativos|learning-resources)/([0-9]+)(?:-[^/?#]+)?\.html", parsed.path)
        prefix = "EP"
    else:
        raise ImportErrorMessage("La URL no corresponde a una plataforma admitida")
    manual = payload.get("manual") is True
    if payload.get("manual") not in (None, True, False) or (manual and payload["verified"]):
        raise ImportErrorMessage("Estado manual inválido")
    if (not match and not (manual and resource in parsed.path and parsed.path != "/")) or (match and match.group(1) != resource):
        raise ImportErrorMessage("El ID del enlace no coincide con el recurso detectado")
    activity_id = f"{prefix}-{resource}"
    slug = slugify(title, resource)
    for activity in catalog["activities"]:
        source = activity.get("source", {})
        if activity["id"] == activity_id or (source.get("platformId"), source.get("resourceId")) == (platform, resource):
            raise ImportErrorMessage("Esta actividad ya existe en el catálogo")
        if url in {source.get("url"), source.get("canonicalUrl"), source.get("originalUrl")}:
            raise ImportErrorMessage("Esta URL ya existe en el catálogo")
        if slug == activity["slug"] or slug in activity.get("legacySlugs", []):
            raise ImportErrorMessage("El slug coincide con una actividad existente")
    return {"id": activity_id, "slug": slug, "url": url, "platform": platform, "resource": resource,
            "title": title, "description": description, "subject": subject, "topics": topics,
            "type": kind, "language": language, "verified": payload["verified"], "sourceTitle": source_title}


def append_with_style(sheet, values):
    row = sheet.max_row + 1
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, value in values.items():
        if name not in columns:
            raise ImportErrorMessage(f"Falta la columna {name} en {sheet.title}")
        cell = sheet.cell(row, columns[name], value)
        previous = sheet.cell(row - 1, columns[name])
        if previous.has_style:
            cell._style = copy(previous._style)
        cell.number_format = previous.number_format
    return row


def import_activity(source, destination, payload, today):
    catalog, _ = generator.load_general_workbook(source)
    data = validate_input(payload, catalog)
    workbook = load_workbook(source)
    stamp = date.fromisoformat(today).isoformat()
    append_with_style(workbook["ACTIVIDADES"], {
        "ID": data["id"], "Slug": data["slug"], "Título": data["title"],
        "Título de origen": data["sourceTitle"], "Descripción": data["description"],
        "ID de asignatura principal": data["subject"], "IDs de asignatura": data["subject"],
        "ID de tema principal": data["topics"][0], "Idioma": data["language"],
        "ID de tipo": data["type"], "Fecha de publicación": stamp, "Fecha de actualización": stamp,
        "Estado editorial": "published", "Clase de origen": "external",
        "ID de plataforma": data["platform"], "ID de recurso": data["resource"],
        "URL pública": data["url"], "URL canónica": data["url"], "URL original": data["url"],
        "Estado del enlace": "verified" if data["verified"] else "unverified",
        "Última verificación": stamp if data["verified"] else None,
        "Título verificado": bool(data["verified"] and data["sourceTitle"]),
        "Fuente del título": "Página pública de la actividad" if data["sourceTitle"] else "Entrada manual",
        "Fuente de la descripción": "Entrada manual",
    })
    for topic in data["topics"]:
        append_with_style(workbook["ACTIVIDAD_TEMA"], {"ID de actividad": data["id"], "ID de tema": topic})
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=destination.parent, delete=False) as temp:
        temporary = Path(temp.name)
    try:
        workbook.save(temporary)
        generated, locations = generator.load_general_workbook(temporary)
        errors, _ = validate_catalog(generated, locations)
        if errors:
            raise ImportErrorMessage("El catálogo no supera la validación: " + "; ".join(e["problem"] for e in errors[:5]))
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return {"id": data["id"], "slug": data["slug"], "url": f"/actividad/{data['slug']}/"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload", type=Path, required=True)
    parser.add_argument("--input", type=Path, default=DEFAULT_BOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_BOOK)
    parser.add_argument("--date", default=date.today().isoformat())
    args = parser.parse_args()
    try:
        payload = json.loads(args.payload.read_text(encoding="utf-8"))
        print(json.dumps(import_activity(args.input, args.output, payload, args.date), ensure_ascii=False))
    except (ValueError, KeyError, TypeError) as error:
        raise SystemExit(str(error)) from error


if __name__ == "__main__":
    main()
