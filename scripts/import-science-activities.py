#!/usr/bin/env python3
"""Import the source workbook into the Site's generated JSON dataset."""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


def clean(value):
    if isinstance(value, str):
        return value.strip()
    return value


def slugify(value: str) -> str:
    text = unicodedata.normalize("NFKD", value)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text


def rows_as_dicts(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in rows[0]]
    return [
        {headers[index]: clean(value) for index, value in enumerate(row)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: import-science-activities.py INPUT.xlsx OUTPUT.json")

    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    workbook = openpyxl.load_workbook(source, data_only=False)

    required_sheets = {
        "ACTIVIDADES",
        "SECCIONES",
        "RELACIONES",
        "CONTROL DE CALIDAD",
        "INFORMACIÓN",
        "AUDITORÍA",
    }
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise SystemExit(f"missing sheets: {sorted(missing)}")

    activity_rows = rows_as_dicts(workbook["ACTIVIDADES"])
    section_rows = rows_as_dicts(workbook["SECCIONES"])
    relation_rows = rows_as_dicts(workbook["RELACIONES"])
    quality_rows = rows_as_dicts(workbook["CONTROL DE CALIDAD"])
    information_rows = rows_as_dicts(workbook["INFORMACIÓN"])
    audit_rows = rows_as_dicts(workbook["AUDITORÍA"])

    sections = []
    sections_by_id = {}
    for row in sorted(section_rows, key=lambda item: item["Orden original"]):
        section = {
            "id": row["ID de sección"],
            "order": int(row["Orden original"]),
            "name": row["Nombre para la web"],
            "slug": slugify(row["Nombre para la web"]),
            "originalName": row["Nombre original del documento"],
            "historicalCourse": row["Curso de origen indicado en el documento"],
            "note": row["Notas"],
            "declaredCount": int(row["Número de actividades únicas relacionadas"]),
        }
        sections.append(section)
        sections_by_id[section["id"]] = section

    section_ids_by_activity = defaultdict(list)
    for row in relation_rows:
        activity_id = row["ID de actividad"]
        section_id = row["ID de sección"]
        if section_id not in sections_by_id:
            raise SystemExit(f"orphan section relation: {section_id}")
        section_ids_by_activity[activity_id].append(section_id)

    activities = []
    activity_ids = set()
    slugs = set()
    urls = set()
    for row in activity_rows:
        activity_id = row["ID de actividad"]
        slug = row["Slug o identificador web"]
        url = row["URL pública"]
        if activity_id in activity_ids:
            raise SystemExit(f"duplicate activity id: {activity_id}")
        if slug in slugs:
            raise SystemExit(f"duplicate activity slug: {slug}")
        if url in urls:
            raise SystemExit(f"duplicate public URL: {url}")
        if activity_id not in section_ids_by_activity:
            raise SystemExit(f"activity without section: {activity_id}")

        activity_ids.add(activity_id)
        slugs.add(slug)
        urls.add(url)
        related_ids = sorted(
            section_ids_by_activity[activity_id],
            key=lambda section_id: sections_by_id[section_id]["order"],
        )
        activities.append(
            {
                "id": activity_id,
                "title": row["Título para la web"],
                "originalTitle": row["Título original"],
                "description": row["Descripción para la web"],
                "platform": row["Plataforma"],
                "url": url,
                "canonicalUrl": row["URL canónica"],
                "language": row["Idioma"],
                "type": row["Tipo de actividad"],
                "course": row["Curso de origen"],
                "subject": row["Asignatura"],
                "tags": row["Etiquetas"],
                "keywords": row["Palabras clave de búsqueda"],
                "slug": slug,
                "status": row["Estado del enlace"],
                "verifiedTitle": row["Título verificado en la página"],
                "verifiedAt": row["Fecha de verificación"].date().isoformat()
                if hasattr(row["Fecha de verificación"], "date")
                else str(row["Fecha de verificación"]),
                "sectionIds": related_ids,
                "primarySectionId": related_ids[0],
            }
        )

    orphan_activity_relations = sorted(set(section_ids_by_activity).difference(activity_ids))
    if orphan_activity_relations:
        raise SystemExit(f"orphan activity relations: {orphan_activity_relations}")

    section_counts = Counter(
        section_id for activity in activities for section_id in activity["sectionIds"]
    )
    for section in sections:
        actual = section_counts[section["id"]]
        if actual != section["declaredCount"]:
            raise SystemExit(
                f"section count mismatch for {section['id']}: {actual} != {section['declaredCount']}"
            )
        section["count"] = actual

    multi_section = [activity for activity in activities if len(activity["sectionIds"]) > 1]
    platform_counts = Counter(activity["platform"] for activity in activities)
    language_counts = Counter(activity["language"] for activity in activities)
    type_counts = Counter(activity["type"] for activity in activities)
    course_counts = Counter(activity["course"] for activity in activities)

    expected = {
        "activities": 65,
        "sections": 11,
        "relations": 69,
        "multiSectionActivities": 4,
    }
    actual = {
        "activities": len(activities),
        "sections": len(sections),
        "relations": len(relation_rows),
        "multiSectionActivities": len(multi_section),
    }
    if actual != expected:
        raise SystemExit(f"audit mismatch: {actual} != {expected}")

    payload = {
        "generatedFrom": source.name,
        "verifiedAt": "2026-09-05",
        "sections": sections,
        "activities": activities,
        "qualityIssues": quality_rows,
        "information": information_rows,
        "audit": audit_rows,
        "summary": {
            **actual,
            "platformCounts": dict(platform_counts),
            "languageCounts": dict(language_counts),
            "typeCounts": dict(type_counts),
            "courseCounts": dict(course_counts),
            "multiSectionIds": [activity["id"] for activity in multi_section],
        },
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(payload["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
